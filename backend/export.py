import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from schemas import Resume


def _contact_line(r: Resume) -> str:
    return "  |  ".join(x for x in [r.email, r.phone, r.location, *r.links] if x)


def _dates(start: str, end: str) -> str:
    return " - ".join(x for x in (start, end) if x)


def safe_filename(name: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "resume"
    return f"{base}_Resume.{ext}"


def to_docx(r: Resume) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(r.name, level=0)
    if contact := _contact_line(r):
        doc.add_paragraph(contact)

    def section(title: str):
        doc.add_heading(title, level=1)

    if r.summary:
        section("Summary")
        doc.add_paragraph(r.summary)
    if r.skills:
        section("Skills")
        doc.add_paragraph(", ".join(r.skills))
    if r.experience:
        section("Experience")
        for e in r.experience:
            p = doc.add_paragraph()
            p.add_run(f"{e.title} - {e.company}").bold = True
            meta = "  |  ".join(x for x in [e.location, _dates(e.start_date, e.end_date)] if x)
            if meta:
                p.add_run(f"\n{meta}").italic = True
            for b in e.bullets:
                doc.add_paragraph(b, style="List Bullet")
    if r.projects:
        section("Projects")
        for pr in r.projects:
            p = doc.add_paragraph()
            p.add_run(pr.name).bold = True
            if pr.description:
                p.add_run(f"\n{pr.description}")
            for b in pr.bullets:
                doc.add_paragraph(b, style="List Bullet")
    if r.education:
        section("Education")
        for ed in r.education:
            p = doc.add_paragraph()
            p.add_run(f"{ed.degree} - {ed.institution}").bold = True
            extra = "  |  ".join(x for x in [ed.year, ed.details] if x)
            if extra:
                p.add_run(f"\n{extra}")
    if r.certifications:
        section("Certifications")
        for c in r.certifications:
            doc.add_paragraph(c, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(r: Resume) -> bytes:
    ss = getSampleStyleSheet()
    name_s = ParagraphStyle("Name", parent=ss["Title"], fontSize=20, spaceAfter=4)
    h_s = ParagraphStyle(
        "H", parent=ss["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=3,
        borderWidth=0, textColor="#222222",
    )
    body = ParagraphStyle("Body", parent=ss["BodyText"], fontSize=10, leading=13)
    meta_s = ParagraphStyle("Meta", parent=body, textColor="#555555", fontName="Helvetica-Oblique")

    def P(text: str, style=body) -> Paragraph:
        return Paragraph(escape(text), style)

    def bullets(items: list[str]):
        return ListFlowable(
            [ListItem(P(b), leftIndent=12) for b in items],
            bulletType="bullet", start="•", leftIndent=12,
        )

    story = [P(r.name, name_s)]
    if contact := _contact_line(r):
        story.append(P(contact, meta_s))

    if r.summary:
        story += [P("SUMMARY", h_s), P(r.summary)]
    if r.skills:
        story += [P("SKILLS", h_s), P(", ".join(r.skills))]
    if r.experience:
        story.append(P("EXPERIENCE", h_s))
        for e in r.experience:
            story.append(Paragraph(f"<b>{escape(e.title)}</b> - {escape(e.company)}", body))
            meta = "  |  ".join(x for x in [e.location, _dates(e.start_date, e.end_date)] if x)
            if meta:
                story.append(P(meta, meta_s))
            if e.bullets:
                story.append(bullets(e.bullets))
            story.append(Spacer(1, 4))
    if r.projects:
        story.append(P("PROJECTS", h_s))
        for pr in r.projects:
            story.append(Paragraph(f"<b>{escape(pr.name)}</b>", body))
            if pr.description:
                story.append(P(pr.description))
            if pr.bullets:
                story.append(bullets(pr.bullets))
            story.append(Spacer(1, 4))
    if r.education:
        story.append(P("EDUCATION", h_s))
        for ed in r.education:
            story.append(Paragraph(f"<b>{escape(ed.degree)}</b> - {escape(ed.institution)}", body))
            extra = "  |  ".join(x for x in [ed.year, ed.details] if x)
            if extra:
                story.append(P(extra, meta_s))
    if r.certifications:
        story += [P("CERTIFICATIONS", h_s), bullets(r.certifications)]

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm, title=f"{r.name} - Resume",
    ).build(story)
    return buf.getvalue()
