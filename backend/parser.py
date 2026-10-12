import io

import docx
import pdfplumber


class ParseError(ValueError):
    pass


def extract_text(data: bytes, ext: str) -> str:
    """Extract plain text from a .pdf or .docx resume."""
    try:
        if ext == ".pdf":
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join((page.extract_text() or "") for page in pdf.pages)
        elif ext == ".docx":
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
        else:
            raise ParseError(f"Unsupported extension: {ext}")
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(
            "Could not read the file. It may be corrupted or password-protected."
        ) from exc

    text = text.strip()
    if len(text) < 50:
        raise ParseError(
            "No readable text found in the resume (scanned image PDFs are not supported)."
        )
    return text
